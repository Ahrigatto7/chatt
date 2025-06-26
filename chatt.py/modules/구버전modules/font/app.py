import streamlit as st
import re
from collections import defaultdict
import pandas as pd
import sqlite3
import os
import json
from fpdf import FPDF
from docx import Document

# === 텍스트 추출 ===
def extract_text(file):
    suffix = file.name.split(".")[-1]
    if suffix == "docx":
        doc = Document(file)
        return "\n".join(p.text for p in doc.paragraphs)
    else:
        return file.read().decode("utf-8")

# === 정교한 해석 규칙 추출기 ===
def extract_rules_advanced(text):
    pattern = re.compile(
        r'(.+?)\s*(경우|일 경우|이라면|이면|일 때|인 경우|했을 때|하게 되면|한다면|할 경우|시|때)\s*[,:]?\s*(.+?)(?:\.|$)'
    )
    rules = {}
    for match in pattern.finditer(text):
        condition = match.group(1).strip()
        connector = match.group(2)
        result = match.group(3).strip()

        condition_clean = clean_clause(condition)
        result_clean = clean_clause(result)

        rules.setdefault(condition_clean, []).append(result_clean)
    return rules

def clean_clause(clause):
    clause = re.sub(r"(은는이가|을를|에|에서|의|로|으로|도|만|과|와|이며|이나|하고|부터|까지)$", "", clause.strip())
    return clause

# === 키워드 분석 ===
def extract_sections(text):
    keywords = ['혼인', '결혼', '再婚', '初婚', '배우자궁', '부궁', '夫星', '妻宮', '副夫宮']
    lines = text.splitlines()
    results = []
    current_block = []
    capture = False

    for line in lines:
        if any(kw in line for kw in keywords):
            capture = True
        if capture:
            if line.strip():
                current_block.append(line.strip())
            else:
                if current_block:
                    results.append("\n".join(current_block))
                    current_block = []
                    capture = False
    return results, keywords

def analyze_keywords(results, keywords):
    keyword_hits = defaultdict(list)
    for text in results:
        for kw in keywords:
            if kw in text:
                keyword_hits[kw].append(text)

    df = {
        "키워드": [],
        "빈도": [],
        "예시 일부": []
    }
    for k, v in keyword_hits.items():
        df["키워드"].append(k)
        df["빈도"].append(len(v))
        df["예시 일부"].append(v[0][:100] + "..." if v else "")
    return pd.DataFrame(df), keyword_hits

# === 저장 ===
def save_text(text, filename):
    with open(filename, "w", encoding="utf-8") as f:
        f.write(text)

def save_json(data, filename):
    with open(filename, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def save_excel(summary_df, keyword_map, filename):
    with pd.ExcelWriter(filename, engine='xlsxwriter') as writer:
        summary_df.to_excel(writer, sheet_name="요약", index=False)
        for kw, texts in keyword_map.items():
            df = pd.DataFrame({kw: texts})
            df.to_excel(writer, sheet_name=kw[:31], index=False)

# === PDF 생성 ===
def generate_pdf_report(filename, rule_data: dict, keyword_summary_df):
    pdf = FPDF()
    pdf.add_page()

    font_path = "NanumGothic.ttf"
    pdf.add_font("Nanum", "", font_path, uni=True)
    pdf.add_font("Nanum", "B", font_path, uni=True)
    pdf.set_font("Nanum", 'B', 16)

    pdf.cell(200, 10, txt="사주 문서 분석 보고서", ln=True, align="C")

    pdf.ln(10)
    pdf.set_font("Nanum", 'B', 12)
    pdf.cell(200, 10, txt="추출된 해석 규칙", ln=True)

    for k, v_list in rule_data.items():
        pdf.set_font("Nanum", 'B', 11)
        pdf.cell(200, 8, txt=f"[{k}]", ln=True)
        pdf.set_font("Nanum", '', 11)
        for v in v_list:
            pdf.multi_cell(0, 6, f"- {v}")

    pdf.add_page()
    pdf.set_font("Nanum", 'B', 12)
    pdf.cell(200, 10, txt="혼인 키워드 분석 요약", ln=True)

    for i, row in keyword_summary_df.iterrows():
        pdf.set_font("Nanum", '', 11)
        line = f"{row['키워드']} ({row['빈도']}회): {row['예시 일부']}"
        pdf.multi_cell(0, 6, line)

    pdf.output(filename)

# === DB 저장 ===
def save_to_db(rules, keyword_map):
    conn = sqlite3.connect("saju_analysis.db")
    cur = conn.cursor()

    cur.execute("""CREATE TABLE IF NOT EXISTS rules (
        id INTEGER PRIMARY KEY AUTOINCREMENT, condition TEXT, result TEXT
    )""")
    cur.execute("""CREATE TABLE IF NOT EXISTS keyword_hits (
        id INTEGER PRIMARY KEY AUTOINCREMENT, keyword TEXT, snippet TEXT
    )""")

    for condition, results in rules.items():
        for result in results:
            cur.execute("INSERT INTO rules (condition, result) VALUES (?, ?)", (condition, result))

    for keyword, snippets in keyword_map.items():
        for snip in snippets:
            cur.execute("INSERT INTO keyword_hits (keyword, snippet) VALUES (?, ?)", (keyword, snip))

    conn.commit()
    conn.close()

# === Streamlit 앱 ===
st.set_page_config(page_title="📄 사주 문서 자동 분석기", layout="wide")
st.title("📄 사주 문서 자동 분석기")

uploaded_file = st.file_uploader("📂 Word (.docx) 또는 텍스트 (.txt) 파일을 업로드하세요", type=["docx", "txt"])

if uploaded_file:
    with st.spinner("파일에서 텍스트 추출 중..."):
        text = extract_text(uploaded_file)
        st.success("텍스트 추출 완료!")
        st.text_area("📑 원문 미리보기", text, height=300)

    if st.button("🔍 분석 시작"):
        with st.spinner("규칙 및 키워드 분석 중..."):
            rules = extract_rules_advanced(text)
            blocks, keywords = extract_sections(text)
            df, keyword_map = analyze_keywords(blocks, keywords)

            os.makedirs("output", exist_ok=True)
            save_text(text, "output/raw.txt")
            save_json(rules, "output/rules.json")
            save_excel(df, keyword_map, "output/keywords.xlsx")
            generate_pdf_report("output/report.pdf", rules, df)
            save_to_db(rules, keyword_map)

        st.success("✅ 분석 및 저장 완료!")
        st.download_button("📥 PDF 보고서 다운로드", open("output/report.pdf", "rb"), file_name="사주_보고서.pdf")
        st.dataframe(df)
